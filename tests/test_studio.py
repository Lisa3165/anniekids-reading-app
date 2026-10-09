import copy
import io
import json
from types import SimpleNamespace
from pathlib import Path
import os
import unittest
from unittest.mock import patch

from docx import Document
from pypdf import PdfWriter

from studio import StudioError, export_docx, extract_material, generate_lesson, get_api_key, split_chapters, validate_result


class StudioTests(unittest.TestCase):
    def setUp(self):
        self.material = extract_material('story.txt', b'Chapter 1\nLily plants a tree. She waters it every day.')
        self.chapters = split_chapters(self.material['text'])
        self.result = {
            'analysis': {d: [{'point': '分析', 'evidence': 'Lily plants a tree.'}] for d in ['What', 'Why', 'How']},
            'chapters': [{'title': 'Chapter 1', 'big_question': 'How does Lily help the tree?',
                          'questions': [{'dimension': d, 'question': 'What does Lily do?', 'answer': 'She plants a tree.', 'evidence': 'Lily plants a tree.'} for d in ['What', 'What', 'Why', 'Why', 'How', 'How']],
                          'task': 'Describe caring for a tree.', 'teacher_notes': '为最终写作积累动作词。'}],
            'final_task': 'Write 60–80 words about caring for a tree.', 'teacher_notes': '核对原文。'}

    def test_txt_encodings_and_words(self):
        self.assertEqual(self.material['words'], 10)
        self.assertEqual(extract_material('a.txt', '标题\nA tree.'.encode('gb18030'))['title'], '标题')
        self.assertEqual(extract_material('a.txt', b'\xef\xbb\xbfA tree.')['text'], 'A tree.')

    def test_docx_paragraph_and_table(self):
        doc = Document()
        doc.add_paragraph('Story')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'A small tree.'
        output = io.BytesIO()
        doc.save(output)
        material = extract_material('a.docx', output.getvalue())
        self.assertIn('A small tree.', material['text'])
        self.assertEqual(material['title'], 'Story')

    def test_bad_and_empty_files(self):
        for name, data in [('x.docx', b'bad'), ('x.pdf', b'bad'), ('x.txt', b' '), ('x.exe', b'hello')]:
            with self.subTest(name=name), self.assertRaises(StudioError):
                extract_material(name, data)

    def test_blank_pdf_rejected(self):
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        output = io.BytesIO()
        writer.write(output)
        with self.assertRaisesRegex(StudioError, 'OCR'):
            extract_material('scan.pdf', output.getvalue())

    def test_text_pdf_extraction(self):
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b'BT /F1 12 Tf 20 250 Td (A small tree.) Tj ET')
        page[NameObject('/Contents')] = writer._add_object(stream)
        output = io.BytesIO()
        writer.write(output)
        self.assertIn('A small tree.', extract_material('story.pdf', output.getvalue())['text'])

    def test_chapters_and_preamble(self):
        chapters = split_chapters('Book title\nChapter One\nFirst story.\nChapter Two\nSecond story.')
        self.assertEqual([c['title'] for c in chapters], ['章节前文字', 'Chapter One', 'Chapter Two'])
        self.assertEqual(len(split_chapters('A story without headings.')), 1)

    def test_validate_and_reject_invented_quote(self):
        self.assertEqual(validate_result(self.result, self.chapters, self.material['text']), self.result)
        altered = copy.deepcopy(self.result)
        altered['chapters'][0]['questions'][0]['evidence'] = 'Lily flies to Mars.'
        with self.assertRaisesRegex(StudioError, '引文'):
            validate_result(altered, self.chapters, self.material['text'])

    def test_reject_missing_dimension_and_changed_chapter(self):
        altered = copy.deepcopy(self.result)
        altered['chapters'][0]['questions'][0]['dimension'] = 'Why'
        with self.assertRaises(StudioError):
            validate_result(altered, self.chapters, self.material['text'])
        altered = copy.deepcopy(self.result)
        altered['chapters'][0]['title'] = 'Invented Chapter'
        with self.assertRaises(StudioError):
            validate_result(altered, self.chapters, self.material['text'])

    def test_no_key(self):
        with patch.dict(os.environ, {}, clear=True), patch('streamlit.secrets', {}), self.assertRaisesRegex(StudioError, 'API Key'):
            generate_lesson(self.material, self.chapters, 'A2')

    def test_key_precedence(self):
        with patch.dict(os.environ, {'LESSON_API_KEY': 'environment-test', 'OPENAI_API_KEY': 'legacy-test'}, clear=True), patch('streamlit.secrets', {'LESSON_API_KEY': 'secrets-test'}):
            self.assertEqual(get_api_key(), 'environment-test')
        with patch.dict(os.environ, {'OPENAI_API_KEY': 'legacy-test'}, clear=True), patch('streamlit.secrets', {'LESSON_API_KEY': 'secrets-test'}):
            self.assertEqual(get_api_key(), 'legacy-test')

    def test_secrets_and_legacy_fallback(self):
        with patch.dict(os.environ, {'LESSON_API_KEY': '  '}, clear=True), patch('streamlit.secrets', {'LESSON_API_KEY': ' secrets-test '}):
            self.assertEqual(get_api_key(), 'secrets-test')
        with patch.dict(os.environ, {}, clear=True), patch('streamlit.secrets', {'OPENAI_API_KEY': 'legacy-secrets-test'}):
            self.assertEqual(get_api_key(), 'legacy-secrets-test')

    def test_missing_secrets_file(self):
        from streamlit.errors import StreamlitSecretNotFoundError
        with patch.dict(os.environ, {}, clear=True), patch('streamlit.secrets') as secrets:
            secrets.get.side_effect = StreamlitSecretNotFoundError('Missing secrets')
            self.assertIsNone(get_api_key())

    def test_secrets_generation(self):
        response = SimpleNamespace(choices=[SimpleNamespace(finish_reason='stop', message=SimpleNamespace(content=json.dumps(self.result)))])
        with patch.dict(os.environ, {}, clear=True), patch('streamlit.secrets', {'LESSON_API_KEY': 'secrets-test'}), patch('openai.OpenAI') as client:
            client.return_value.chat.completions.create.return_value = response
            self.assertEqual(generate_lesson(self.material, self.chapters, 'A2'), self.result)
            self.assertEqual(client.call_args.kwargs['api_key'], 'secrets-test')

    def test_secrets_ui_does_not_display_key(self):
        from streamlit.testing.v1 import AppTest
        with patch.dict(os.environ, {}, clear=True), patch('streamlit.secrets', {'LESSON_API_KEY': 'secrets-test'}):
            app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / 'app.py')).run(timeout=20)
        self.assertEqual(len(app.exception), 0)
        self.assertTrue(any('已检测到 API Key' in item.value for item in app.success))
        for kind in ('text', 'markdown', 'code', 'info', 'success', 'error', 'warning'):
            self.assertFalse(any('secrets-test' in str(item.value) for item in app.get(kind)))

    def test_word_export_is_editable(self):
        output = export_docx(self.material, self.result, 'A2')
        doc = Document(io.BytesIO(output))
        self.assertEqual(len(doc.tables), 4)
        self.assertEqual(doc.tables[-1].rows[1].cells[1].text, 'What does Lily do?')
        self.assertTrue(any(p.text == self.result['final_task'] for p in doc.paragraphs))

    def test_streamlit_initial_view(self):
        from streamlit.testing.v1 import AppTest
        with patch.dict(os.environ, {'LESSON_API_KEY': '', 'OPENAI_API_KEY': ''}), patch('streamlit.secrets', {}):
            app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / 'app.py')).run(timeout=20)
        self.assertEqual(len(app.exception), 0)
        self.assertTrue(any('未配置 API Key' in item.value for item in app.info))
        self.assertEqual(app.selectbox[0].options, ['A2', 'B1'])

    def test_ai_request_and_validation(self):
        response = SimpleNamespace(choices=[SimpleNamespace(finish_reason='stop', message=SimpleNamespace(content=json.dumps(self.result)))])
        with patch.dict(os.environ, {'LESSON_API_KEY': 'test-key'}), patch('openai.OpenAI') as client:
            client.return_value.chat.completions.create.return_value = response
            result = generate_lesson(self.material, self.chapters, 'B1')
            request = client.return_value.chat.completions.create.call_args.kwargs
            self.assertEqual(json.loads(request['messages'][1]['content'])['level'], 'B1')
            self.assertEqual(result, self.result)

    def test_ai_truncation_rejected(self):
        response = SimpleNamespace(choices=[SimpleNamespace(finish_reason='length')])
        with patch.dict(os.environ, {'LESSON_API_KEY': 'test-key'}), patch('openai.OpenAI') as client:
            client.return_value.chat.completions.create.return_value = response
            with self.assertRaisesRegex(StudioError, '完整'):
                generate_lesson(self.material, self.chapters, 'A2')


if __name__ == '__main__':
    unittest.main()
