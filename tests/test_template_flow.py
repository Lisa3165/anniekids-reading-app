import io
import unittest
import zipfile
from unittest.mock import patch
from docx import Document
from template_flow import analyze_book, make_lessons, validate_lesson, export_pack
from studio import StudioError

class TemplateTests(unittest.TestCase):
    def sample(self):
        return dict(title='Lesson', activities=[dict(minutes=80, activity='Read A')], teacher='教师', worksheet='Question A\n________', answers='Answer A')

    def test_large_text_complete_and_bounded(self):
        text = ('A tree. ' * 8000)
        with patch('template_flow.generate_lesson', return_value={}) as generate:
            notes = analyze_book({'title':'Book'}, [{'title':'Chapter 1', 'text':text}], 'A2', lambda *a: None)
        self.assertGreater(len(notes), 1)
        pieces = [c.args[1][0]['text'] for c in generate.call_args_list]
        self.assertEqual(''.join(pieces), text)
        self.assertTrue(all(len(p) <= 18000 for p in pieces))

    def test_wrong_timing_rejected(self):
        with self.assertRaises(StudioError): validate_lesson(self.sample(), 40)

    def test_pack_three_editable_documents(self):
        with zipfile.ZipFile(io.BytesIO(export_pack('Framework', [self.sample()]))) as z:
            self.assertEqual(len(z.namelist()), 3)
            worksheet = Document(io.BytesIO(z.read('学生作业纸.docx')))
            self.assertTrue(any('Question A' in p.text for p in worksheet.paragraphs))
            self.assertFalse(any('Answer A' in p.text for p in worksheet.paragraphs))

    def test_generation_count_and_previous_context(self):
        with patch('template_flow.request_json', side_effect=[self.sample(), self.sample()]) as request:
            lessons = make_lessons('Plan', {'lessons':2,'minutes':80}, lambda *a: None)
            self.assertEqual(len(lessons), 2)
            self.assertEqual(request.call_args.args[1]['previous_lessons'], ['Lesson'])
